"""IPPT monitoring report parsing and derivation (pure, no DB).

Implements the ingestion contract of docs/IPPT_MONITORING.md §2–§4 and
§6.2/§6.4:

- filenames ``{REPORT}_{STATE}_{YYYYMMDD}.csv`` (or ``.xlsx``) →
  (file_kind, report_date)
- header-name-based column mapping — column order is not stable across
  files (``IPPT_FAILED`` and ``NSFIT_IN_PROGRESS`` share a schema in
  different orders, §2.2), so parsers must never map by position
- ``-`` is the null/NA sentinel in every column → ``None``
- the unified ``state`` derivation is a pure function of
  ``(file_kind, ippt_status, fit_sessions)`` — completion is decided by
  designation/file membership, never by counting FIT sessions (§2.3)
- cheap per-row validation guards quarantine bad rows instead of
  rejecting the file (§4)
- window date derivation, floored calendar-month arithmetic, and the
  3/6/9-month escalation-tier rules (§6.4)

Everything here is synchronous and side-effect-free so the unit tests
can pin the report contract exactly.
"""

from __future__ import annotations

import csv
import datetime
import io
import re
from dataclasses import dataclass, field
from datetime import date, timedelta

from openpyxl import load_workbook

from parade_state.models.ippt import FILE_KINDS

#: Report statuses that close the window's obligation via an award
#: (display order lowest → highest, §1.2).
AWARD_STATUSES: tuple[str, ...] = (
    "Pass",
    "Pass with Incentive",
    "Silver",
    "Gold",
)

#: States that mean the window's obligation is closed — no tier ever
#: highlights these ("waived counts as fulfilled in all tiers", §6.4).
FULFILLED_STATES: tuple[str, ...] = (
    "met_award",
    "met_voluntary_fit",
    "met_mandatory_fit",
    "waived",
)

_FILENAME_RE = re.compile(
    r"^(IPPT|NSFIT)_(COMPLETED|FAILED|NOT_ATTEMPTED|IN_PROGRESS|NOT_STARTED)_"
    r"(\d{4})(\d{2})(\d{2})\.(csv|xlsx)$",
    re.IGNORECASE,
)

_STATE_BY_TOKEN = {
    "IPPT_COMPLETED": "ippt_completed",
    "IPPT_FAILED": "ippt_failed",
    "IPPT_NOT_ATTEMPTED": "ippt_not_attempted",
    "NSFIT_COMPLETED": "nsfit_completed",
    "NSFIT_IN_PROGRESS": "nsfit_in_progress",
    "NSFIT_NOT_STARTED": "nsfit_not_started",
}

# Column layouts by family (§2.2). Two files share the same fields in a
# different order, so mapping is by header name, never by position.
_FAMILY_HEADERS: dict[str, tuple[str, ...]] = {
    "ippt_completed": (
        "Rank",
        "Name",
        "Unit",
        "Sub-unit",
        "IPPT status",
        "FIT sessions completed",
    ),
    "nsfit_completed": (
        "Rank",
        "Name",
        "Unit",
        "Sub-unit",
        "IPPT status",
        "FIT sessions completed",
    ),
    "ippt_failed": (
        "Rank",
        "Name",
        "Unit",
        "Sub-unit",
        "Window close",
        "IPPT status",
        "FIT sessions completed",
        "Reminder sent",
        "Last sent",
    ),
    "nsfit_in_progress": (
        "Rank",
        "Name",
        "Unit",
        "Sub-unit",
        "Window close",
        "IPPT status",
        "FIT sessions completed",
        "Reminder sent",
        "Last sent",
    ),
    "ippt_not_attempted": (
        "Rank",
        "Name",
        "Unit",
        "Sub-unit",
        "Window close",
        "Booked Status",
        "FFI",
        "Reminder sent",
        "Last sent",
    ),
    "nsfit_not_started": (
        "Rank",
        "Name",
        "Unit",
        "Sub-unit",
        "Window close",
        "Booked Status",
        "FFI",
        "Reminder sent",
        "Last sent",
    ),
}


class FilenameError(ValueError):
    """A report filename does not match
    ``{REPORT}_{STATE}_{YYYYMMDD}.csv`` (or ``.xlsx``)."""


class ReportFileError(ValueError):
    """A report file is unusable at file level (unreadable content,
    empty file, or missing required columns for its family). File-level
    problems reject the whole submission — per-row problems quarantine
    single rows (§4)."""


def parse_filename(filename: str) -> tuple[str, date]:
    """Extract ``(file_kind, report_date)`` from a report filename.

    Raises:
        FilenameError: when the name does not match the canonical
            ``{REPORT}_{STATE}_{YYYYMMDD}.csv``/``.xlsx`` pattern.
    """
    name = filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    match = _FILENAME_RE.match(name)
    if not match:
        raise FilenameError(
            f"Report filename {filename!r} does not match the expected "
            f"{'{REPORT}_{STATE}_{YYYYMMDD}.csv|.xlsx'} pattern"
        )
    token = f"{match.group(1)}_{match.group(2)}".upper()
    report_date = date(int(match.group(3)), int(match.group(4)), int(match.group(5)))
    return _STATE_BY_TOKEN[token], report_date


def normalize_name(name: str) -> str:
    """Normalize a full name for matching (§3.2): uppercase, collapsed
    internal whitespace. Commas, ``S/O`` and apostrophes are part of the
    name and are kept."""
    return " ".join(name.split()).upper()


@dataclass
class IpptReportRow:
    """One parsed report row; ``None`` = ``-`` sentinel or absent column."""

    line_number: int
    rank: str
    full_name: str
    unit: str
    sub_unit: str
    state: str
    window_close_days: int | None = None
    ippt_status: str | None = None
    fit_sessions: int | None = None
    booked: bool | None = None
    ffi: str | None = None
    reminder_sent: str | None = None
    last_sent: str | None = None


@dataclass
class QuarantinedRow:
    """A rejected row kept for review (§4): raw line + why."""

    line_number: int
    raw_row: str
    reason: str


@dataclass
class ParsedReportFile:
    file_kind: str
    filename: str
    report_date: date
    rows: list[IpptReportRow] = field(default_factory=list)
    quarantined: list[QuarantinedRow] = field(default_factory=list)


def _clean(value: str | None) -> str | None:
    """Strip a cell; ``-`` is the null sentinel (§2.3) → ``None``."""
    if value is None:
        return None
    stripped = value.strip()
    if stripped == "-":
        return None
    return stripped


def _parse_int(value: str | None) -> int | None:
    return None if value is None else int(value)


def _parse_bool(value: str | None) -> bool | None:
    if value is None:
        return None
    folded = value.casefold()
    if folded == "yes":
        return True
    if folded == "no":
        return False
    raise ValueError(f"Booked Status must be Yes or No, got {value!r}")


def _parse_ffi(value: str | None) -> str | None:
    if value is None:
        return None
    folded = value.casefold()
    if folded in ("na", "fit", "pending"):
        return folded
    raise ValueError(f"FFI must be NA, Fit or Pending, got {value!r}")


def derive_state(
    file_kind: str, ippt_status: str | None, fit_sessions: int | None
) -> str | None:
    """Map ``(file_kind, ippt_status, fit_sessions)`` to the unified
    state (§6.2). ``None`` = unrecognized combination → the row is
    quarantined for review rather than guessed (§4: when in doubt, flag).

    Completion comes only from designations: ``*_COMPLETED`` membership
    is authoritative, and FIT session counts never decide completion —
    ``Complete FIT`` rows legitimately show 9 sessions.
    """
    if file_kind == "nsfit_completed":
        # Discharged by 10 sessions or a mid-programme pass; no
        # 10-session row exists in the sample (§5), so membership is the
        # signal, whatever the status column shows.
        return "met_mandatory_fit"
    if file_kind == "ippt_completed":
        if ippt_status in AWARD_STATUSES:
            return "met_award"
        if ippt_status == "Complete FIT":
            return "met_voluntary_fit"
        if ippt_status is None:
            return "waived"
        return None
    if file_kind == "ippt_failed":
        if ippt_status == "Fail":
            return "failed_can_reattempt"
        if ippt_status is None and fit_sessions is not None and 1 <= fit_sessions <= 9:
            return "voluntary_fit_in_progress"
        return None
    if file_kind == "nsfit_in_progress":
        if ippt_status == "Fail":
            return "mandatory_fit_failed_attempt"
        if ippt_status is None and fit_sessions is not None and fit_sessions >= 1:
            return "mandatory_fit_in_progress"
        return None
    if file_kind == "ippt_not_attempted":
        return "ippt_not_attempted"
    if file_kind == "nsfit_not_started":
        return "nsfit_not_started"
    return None


def parse_report_file(filename: str, content: bytes | str) -> ParsedReportFile:
    """Parse one report file into rows + quarantine entries.

    ``.xlsx`` submissions are read from their first worksheet with every
    cell stringified (integers render without a ``.0``); ``.csv``
    submissions are UTF-8/ASCII text. Everything after the raw row
    matrix is format-agnostic. Filename and header problems raise (the
    whole submission is rejected); per-row value/derivation problems
    quarantine the single row and keep the rest of the file (§4).

    Raises:
        FilenameError: unparsable filename.
        ReportFileError: unreadable or empty file, or missing required
            headers for the file's family.
    """
    file_kind, report_date = parse_filename(filename)
    if filename.lower().endswith(".xlsx"):
        rows = _xlsx_rows(filename, content)
    else:
        rows = _csv_rows(filename, content)
    if not rows:
        raise ReportFileError(f"{filename}: file is empty")

    header = rows[0]
    names = [name.strip() for name in header]
    index_by_name: dict[str, int] = {}
    for index, name in enumerate(names):
        if name and name not in index_by_name:  # first occurrence wins
            index_by_name[name] = index

    family = _HEADER_NAMES[file_kind]
    missing = [name for name in family if name not in index_by_name]
    if missing:
        raise ReportFileError(
            f"{filename}: missing required column(s): {', '.join(missing)}"
        )

    def cell(row: list[str], name: str) -> str | None:
        index = index_by_name[name]
        return row[index] if index < len(row) else None

    parsed = ParsedReportFile(
        file_kind=file_kind, filename=filename, report_date=report_date
    )

    for line_number, row in enumerate(rows[1:], start=2):  # line 1 is the header
        if not any((value or "").strip() for value in row):
            continue  # blank/trailing line
        raw = ",".join(value or "" for value in row)
        try:
            result = _parse_row(parsed, row, cell, family, line_number, raw)
        except (ValueError, IndexError) as exc:
            parsed.quarantined.append(
                QuarantinedRow(line_number=line_number, raw_row=raw, reason=str(exc))
            )
            continue
        if result is not None:
            parsed.quarantined.append(result)
    return parsed


def _csv_rows(filename: str, content: bytes | str) -> list[list[str]]:
    """The raw CSV row matrix (row 1 = header), UTF-8/ASCII only."""
    if isinstance(content, bytes):
        try:
            content = content.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ReportFileError(f"{filename}: not UTF-8/ASCII text: {exc}") from exc
    return list(csv.reader(io.StringIO(content)))


def _xlsx_rows(filename: str, content: bytes | str) -> list[list[str]]:
    """The first worksheet as a raw row matrix (row 1 = header), every
    cell stringified like its CSV export would read. Excel's own empty
    cells become ``""`` — the explicit ``-`` sentinel stays the only way
    to write null, matching the CSV contract."""
    try:
        workbook = load_workbook(
            io.BytesIO(content),  # type: ignore[arg-type]
            read_only=True,
            data_only=True,  # formula cells → last computed value
        )
    except Exception as exc:  # noqa: BLE001 — any zip/XML defect is "unusable"
        raise ReportFileError(f"{filename}: unreadable .xlsx workbook: {exc}") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise ReportFileError(f"{filename}: workbook has no sheets")
        return [
            [_cell_text(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()


def _cell_text(value) -> str:
    """Stringify one worksheet cell the way the CSV reader would have."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores 6 as 6.0
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value)


def _parse_row(
    parsed: ParsedReportFile,
    row: list[str],
    cell,
    family: frozenset[str],
    line_number: int,
    raw: str,
) -> QuarantinedRow | None:
    """Parse one CSV row into ``parsed.rows``; return a quarantine entry
    instead of a row when validation or state derivation fails."""

    def has(name: str) -> bool:
        return name in family

    problems: list[str] = []

    rank = _clean(cell(row, "Rank"))
    full_name = _clean(cell(row, "Name"))
    unit = _clean(cell(row, "Unit"))
    sub_unit = _clean(cell(row, "Sub-unit"))
    if not rank:
        problems.append("missing Rank")
    if not full_name:
        problems.append("missing Name")
    if not unit:
        problems.append("missing Unit")
    if not sub_unit:
        problems.append("missing Sub-unit")

    window_close_days: int | None = None
    if has("Window close"):
        raw_close = _clean(cell(row, "Window close"))
        if raw_close is not None:
            try:
                window_close_days = _parse_int(raw_close)
            except ValueError:
                problems.append(f"Window close {raw_close!r} is not an integer")
            else:
                if not 0 <= window_close_days <= 366:
                    problems.append(
                        f"Window close {window_close_days} outside [0, 366]"
                    )

    ippt_status = _clean(cell(row, "IPPT status")) if has("IPPT status") else None

    fit_sessions: int | None = None
    if has("FIT sessions completed"):
        raw_fit = _clean(cell(row, "FIT sessions completed"))
        if raw_fit is not None:
            try:
                fit_sessions = _parse_int(raw_fit)
            except ValueError:
                problems.append(f"FIT sessions completed {raw_fit!r} is not an integer")
            else:
                if not 0 <= fit_sessions <= 10:
                    problems.append(
                        f"FIT sessions completed {fit_sessions} outside [0, 10]"
                    )

    booked: bool | None = None
    if has("Booked Status"):
        raw_booked = _clean(cell(row, "Booked Status"))
        try:
            booked = _parse_bool(raw_booked)
        except ValueError as exc:
            problems.append(str(exc))

    ffi: str | None = None
    if has("FFI"):
        raw_ffi = _clean(cell(row, "FFI"))
        try:
            ffi = _parse_ffi(raw_ffi)
        except ValueError as exc:
            problems.append(str(exc))

    if problems:
        return QuarantinedRow(
            line_number=line_number, raw_row=raw, reason="; ".join(problems)
        )

    state = derive_state(parsed.file_kind, ippt_status, fit_sessions)
    if state is None:
        status_display = ippt_status if ippt_status is not None else "-"
        fit_display = str(fit_sessions) if fit_sessions is not None else "-"
        return QuarantinedRow(
            line_number=line_number,
            raw_row=raw,
            reason=(
                f"unrecognized {parsed.file_kind} row: "
                f"IPPT status {status_display!r}, FIT sessions {fit_display!r}"
            ),
        )

    parsed.rows.append(
        IpptReportRow(
            line_number=line_number,
            rank=rank,
            full_name=full_name,
            unit=unit,
            sub_unit=sub_unit,
            state=state,
            window_close_days=window_close_days,
            ippt_status=ippt_status,
            fit_sessions=fit_sessions,
            booked=booked,
            ffi=ffi,
            reminder_sent=(
                _clean(cell(row, "Reminder sent")) if has("Reminder sent") else None
            ),
            last_sent=_clean(cell(row, "Last sent")) if has("Last sent") else None,
        )
    )
    return None


def window_dates(report_date: date, window_close_days: int) -> tuple[date, date]:
    """Derive ``(window_start, window_end)`` from an outstanding row.

    ``window_end`` = report_date + window_close days; ``window_start`` =
    the birthday = (window_end + 1 day) − 1 year (§6.1). Feb-29 birthdays
    settle on Feb-28 in common years.
    """
    window_end = report_date + timedelta(days=window_close_days)
    day_after = window_end + timedelta(days=1)
    try:
        window_start = day_after.replace(year=day_after.year - 1)
    except ValueError:  # Feb 29 in a non-leap target year
        window_start = day_after.replace(year=day_after.year - 1, day=28)
    return window_start, window_end


def months_between(window_start: date, as_of: date) -> int:
    """Floored calendar-month difference (§6.4): whole months from
    ``window_start`` to ``as_of``, minus one when the day-of-month
    hasn't been reached yet. Clamped at zero."""
    months = 12 * (as_of.year - window_start.year) + (as_of.month - window_start.month)
    if as_of.day < window_start.day:
        months -= 1
    return max(months, 0)


def tier_band(months_to_close: int) -> str | None:
    """The escalation band for a months-to-close value, or ``None`` while
    more than 9 months remain before the window closes.

    Bands are named for proximity to the close (the unit's nomenclature):
    ``tier_3`` — at most 3 months before window close — is the most
    urgent, ``tier_6`` the next, ``tier_9`` the least. Under a ~12-month
    window this partitions the timeline identically to the §6.4
    months-elapsed bands (elapsed ≥ 9 ⇔ ≤ 3 before close), relabelled
    2026-10-08.
    """
    if months_to_close <= 3:
        return "tier_3"
    if months_to_close <= 6:
        return "tier_6"
    if months_to_close <= 9:
        return "tier_9"
    return None


TIER_LABELS: dict[str, str] = {
    "tier_3": "3 months before window close",
    "tier_6": "6 months before window close",
    "tier_9": "9 months before window close",
}


def tier_highlighted(
    band: str | None,
    state: str,
    booked: bool | None,
    ffi: str | None,
) -> bool:
    """Whether the per-tier highlight list includes this observation.

    Governed by "when in doubt, highlight" (§6.4) — false positives can
    be relaxed later, false negatives go unnoticed. The criteria attach
    to the window stage, which the band names express as months before
    the window closes:

    - tier_3 (most urgent): everything unfulfilled is highlighted —
      the window is about to close.
    - tier_6: real progress or its exact equivalent is required — an
      attempt, any fulfilled/waived state, or ``ippt_not_attempted``
      booked *after* screening cleared (``fit``). Incomplete FIT routes
      are chased, mandatory included.
    - tier_9 (least urgent): only the not-yet-started highlight —
      ``ippt_not_attempted`` that hasn't booked, or booked with the 35+
      screening still ``pending`` (screening credit requires booking:
      ``fit`` + booked).
    """
    if state in FULFILLED_STATES:
        return False
    if band is None:
        return False
    if band == "tier_3":
        return True  # window about to close: everything unfulfilled
    if band == "tier_6":
        if state == "failed_can_reattempt":
            return False
        if state == "ippt_not_attempted":
            return not (booked and ffi == "fit")
        # voluntary/mandatory FIT in progress, nsfit_not_started
        return True
    # tier_9 (least urgent)
    if state != "ippt_not_attempted":
        return False
    if not booked:
        return True
    return ffi == "pending"


#: The exact header-name set each file kind is parsed with (used by the
#: row parser to know which columns a family carries).
_HEADER_NAMES: dict[str, frozenset[str]] = {
    kind: frozenset(headers) for kind, headers in _FAMILY_HEADERS.items()
}

assert set(_HEADER_NAMES) == set(FILE_KINDS)
