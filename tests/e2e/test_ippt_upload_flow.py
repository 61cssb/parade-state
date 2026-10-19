"""E2E: the IPPT upload → dashboard → window flow in a real browser.

Drives system Chrome (via Playwright) against a real uvicorn server on a
freshly migrated + seeded database. The canonical fixtures upload covers
the full client-side validation (six canonical files) and the server's
atomic ingest; the dashboard and window assertions pin what the feature
is for: seeing who needs chasing. The fixtures may hold several monthly
snapshots — each report date uploads as its own six-file submission,
oldest first (2026-10-08 hardening: order never matters, the recompute
pass links backfilled months).

The database is rebuilt per test (see conftest), so repeated runs and
re-uploads always start from the same state.
"""

import io
from pathlib import Path

import pytest

pytestmark = pytest.mark.e2e

FIXTURES = Path(__file__).resolve().parent.parent.parent / "fixtures" / "ippt"


def _fixture_snapshots() -> list[list[str]]:
    """Canonical report files grouped into six-file snapshots, one per
    report date, oldest date first."""
    by_date: dict[str, list[str]] = {}
    for path in sorted(FIXTURES.glob("*.csv")):
        by_date.setdefault(path.stem.rsplit("_", 1)[1], []).append(str(path))
    return [files for _, files in sorted(by_date.items())]


def _upload(page, base_url: str, files: list[str]) -> None:
    """Upload files through the real UI and wait for the request to finish.

    The upload button is disabled while the request is in flight and
    re-enabled in a ``finally`` — a reliable completion signal.
    """
    page.goto(f"{base_url}/ippt/upload")
    page.set_input_files("#ipptFiles", files)
    page.click("#uploadBtn")
    page.wait_for_function("document.querySelector('#uploadBtn').disabled === false")


def test_client_side_rejects_incomplete_selection(page, e2e_server):
    snapshots = _fixture_snapshots()
    if not snapshots:
        pytest.skip("canonical IPPT fixtures not present — /fixtures is gitignored")

    page.goto(f"{e2e_server}/ippt/upload")
    page.set_input_files("#ipptFiles", snapshots[0][:2])
    page.click("#uploadBtn")
    status = page.inner_text("#uploadStatus")
    assert "exactly the six report files" in status


def test_upload_fixtures_and_read_dashboard(page, e2e_server):
    snapshots = _fixture_snapshots()
    if not snapshots:
        pytest.skip("canonical IPPT fixtures not present — /fixtures is gitignored")

    # One six-file submission per report date, oldest first.
    for index, files in enumerate(snapshots):
        _upload(page, e2e_server, files)
        status = page.inner_text("#uploadStatus")
        if index == 0:
            assert "2026-09-11 ingested" in status
            result = page.inner_text("#uploadResult")
            assert "284 rows" in result
            assert "0 quarantined" in result
        else:
            assert "ingested" in status

    # The dashboard reflects the ingest (most-urgent tier tab by default).
    page.click("text=Open the dashboard")
    page.wait_for_selector("text=Snapshot history")
    body = page.content()
    assert "LEW EE KENT" in body  # the most urgent failure (4 days left)
    assert "3 months before window close" in body
    assert "9 months before window close" in body

    # Through to a per-person window page via a roster link.
    page.click("a:has-text('LEW EE KENT')")
    page.wait_for_selector("text=Trajectory")
    detail = page.content()
    assert "Attempted — failed (can re-attempt)" in detail
    assert "Health screening (FFI)" in detail


def test_reupload_replaces_same_report_date(page, e2e_server):
    snapshots = _fixture_snapshots()
    if not snapshots:
        pytest.skip("canonical IPPT fixtures not present — /fixtures is gitignored")

    first = snapshots[0]
    _upload(page, e2e_server, first)
    assert "ingested" in page.inner_text("#uploadStatus")

    # Same six files again: replace-on-re-ingest keeps the dashboard stable.
    _upload(page, e2e_server, first)
    assert "replaced the previous ingest" in page.inner_text("#uploadStatus")

    page.goto(f"{e2e_server}/ippt/dashboard")
    page.wait_for_selector("text=Snapshot history")
    assert "284" in page.content()


def test_remove_from_tracking_flow(page, e2e_server):
    """The decided 2026-10-08 remove-from-tracking: per-row super-admin
    button → reason dialog → row vanishes from the dashboard; the upload
    page lists him for re-inclusion, which restores him."""
    snapshots = _fixture_snapshots()
    if not snapshots:
        pytest.skip("canonical IPPT fixtures not present — /fixtures is gitignored")

    _upload(page, e2e_server, snapshots[0])

    # The most urgent row (LEW EE KENT, 4 days left) sorts first.
    page.goto(f"{e2e_server}/ippt/dashboard")
    page.wait_for_selector("text=LEW EE KENT")
    page.click(".ippt-action >> nth=0")
    page.wait_for_selector("#excludeDialog[open]")
    name = page.inner_text("#excludeName").strip()
    assert name
    page.fill("#excludeReason", "E2E exclusion check")
    page.click("#excludeSubmit")
    page.wait_for_load_state("load")  # the handler reloads on success
    page.wait_for_selector("text=1 removed from tracking")
    body = page.content()
    assert name not in body  # hidden everywhere on the dashboard

    # The upload page keeps him re-includable, with his reason.
    page.goto(f"{e2e_server}/ippt/upload")
    page.wait_for_selector("#excludedCard >> text=Removed from tracking")
    assert name in page.inner_text("#excludedCard")
    assert "E2E exclusion check" in page.inner_text("#excludedCard")
    page.click("#excludedCard .ippt-action")
    page.wait_for_selector("text=Everyone seen in the reports is currently tracked.")

    # Re-included: back on the dashboard, and the counter is gone (the
    # status line only mentions removals while any exist).
    page.goto(f"{e2e_server}/ippt/dashboard")
    page.wait_for_selector(f"text={name}")
    assert "removed from tracking" not in page.content()


def test_reinclude_from_upload_result_panel(page, e2e_server):
    """Exclude, then re-upload the same date: the ingest result panel must
    offer re-inclusion, and its button must actually re-include (the
    panel button carries the serviceman id — regression for the silent
    '/servicemen/undefined' fetch)."""
    snapshots = _fixture_snapshots()
    if not snapshots:
        pytest.skip("canonical IPPT fixtures not present — /fixtures is gitignored")

    _upload(page, e2e_server, snapshots[0])
    page.evaluate(
        """async () => {
            const dash = await fetch('/api/v1/ippt/dashboard').then(r => r.json());
            const target = dash.servicemen.find(e => e.full_name === 'LEW EE KENT');
            const response = await fetch(
                '/api/v1/ippt/servicemen/' + target.serviceman_id + '/exclusion',
                {
                    method: 'POST',
                    credentials: 'same-origin',
                    headers: {'Content-Type': 'application/json'},
                    body: JSON.stringify({reason: 'panel regression test'}),
                },
            );
            if (!response.ok) throw new Error('exclude failed: ' + response.status);
        }"""
    )

    _upload(page, e2e_server, snapshots[0])  # re-ingest surfaces him again
    page.wait_for_selector("#uploadResult li[data-reappeared] button")
    assert "panel regression test" in page.inner_text("#uploadResult")
    page.click("#uploadResult li[data-reappeared] button")
    page.wait_for_selector("#uploadResult li[data-reappeared]", state="detached")

    page.goto(f"{e2e_server}/ippt/dashboard")
    page.wait_for_selector("text=LEW EE KENT")
    assert "removed from tracking" not in page.content()


def test_export_xlsx_button_downloads_full_workbook(page, e2e_server):
    """The Export XLSX button (decided 2026-10-19) downloads a five-tab
    workbook of the FULL dataset: real browser download, dashboard tab
    order, and — with a state filter active on screen — the export still
    carries the servicemen the filter hides."""
    from datetime import datetime

    from openpyxl import load_workbook

    snapshots = _fixture_snapshots()
    if not snapshots:
        pytest.skip("canonical IPPT fixtures not present — /fixtures is gitignored")

    _upload(page, e2e_server, snapshots[0])

    # A state filter narrows the on-screen tier view...
    page.goto(f"{e2e_server}/ippt/dashboard?tier=tier_3&state=ippt_not_attempted")
    page.wait_for_selector("text=Export XLSX")
    assert "of" in page.inner_text(".ippt-section-title")  # "N of M" badge

    with page.expect_download() as download_info:
        page.click("text=Export XLSX")
    download = download_info.value
    assert download.suggested_filename.startswith("ippt_dashboard_")

    # openpyxl needs a real extension or a stream — the download lands in
    # an extensionless temp path, so hand it the bytes.
    book = load_workbook(io.BytesIO(Path(download.path()).read_bytes()))
    assert book.sheetnames == [
        "3 months",
        "6 months",
        "9 months",
        "On track",
        "Fulfilled",
    ]

    tier3 = book["3 months"]
    assert [cell.value for cell in tier3[1]][:4] == [
        "Rank",
        "Name",
        "Sub-unit",
        "State",
    ]
    names = [row[1] for row in tier3.iter_rows(min_row=2, values_only=True)]
    assert "LEW EE KENT" in names  # ...but the export ignores the filter

    # "Window ends" is a real date cell; "Days left" a real number.
    first = next(tier3.iter_rows(min_row=2, values_only=True))
    assert isinstance(first[6], datetime)
    assert isinstance(first[7], int)
